"""Focused offline regressions for Office provenance and extraction boundaries."""

import hashlib
import io
import json
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
import openpyxl
from openpyxl.worksheet.table import Table
from collect_source_originals import read_jsonl, save_object
from merge_evidence_batch import prepare
from validate_source_originals import check_artifacts
import recover_preserved_office_sources as recovery


TEXT = "Retain the original evidence and preserve exact values, formulas and source identity for independent legal review."


def workbook_bytes(empty=False):
    book = openpyxl.Workbook()
    book.active.title = "Reporting"
    if not empty:
        sheet = book.active
        sheet["A1"] = TEXT
        sheet["A2"] = "Metric"
        sheet["B2"] = "Value"
        sheet["A3"] = "Stored zero"
        sheet["B3"] = 0
        sheet["A4"] = "Percentage"
        sheet["B4"] = 0.125
        sheet["B4"].number_format = "0.0%"
        sheet["C3"] = "=SUM(B3:B4)"
        sheet["D3"] = "=1+2"
        sheet["E3"] = "=0"
        sheet.merge_cells("A6:B6")
        sheet["A6"] = "Merged source label"
        sheet.add_table(Table(displayName="SourceTable", ref="A2:B4"))
        sheet.row_dimensions[4].hidden = True
        sheet.column_dimensions["F"].hidden = True
        hidden = book.create_sheet("Hidden source")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "Hidden source material must remain available in recovered research text."
    buffer = io.BytesIO()
    book.save(buffer)
    if empty:
        return buffer.getvalue()
    rewritten = io.BytesIO()
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as dest:
        for part in source.infolist():
            raw = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                raw = raw.replace(b"<f>SUM(B3:B4)</f><v></v>", b"<f>SUM(B3:B4)</f><v>0.125</v>")
                raw = raw.replace(b"<f>0</f><v></v>", b"<f>0</f><v>0</v>")
            dest.writestr(part, raw)
    return rewritten.getvalue()


class OfficeRecoveryTests(unittest.TestCase):
    def setUp(self):
        base = recovery.STAGE / "office-recovery/work"
        base.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(prefix="test-office-", dir=base)
        self.addCleanup(self.temp.cleanup)
        self.archive = Path(self.temp.name) / "archive"
        self.output = Path(self.temp.name) / "output"
        self.archive.mkdir()
        self.addCleanup(patch.stopall)
        patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden")).start()
        patch.object(socket, "create_connection", side_effect=AssertionError("Network forbidden")).start()

    def prior(self, data, url="https://example.gov.au/report.xlsx.aspx", **changes):
        relative, digest = save_object(self.archive, data, ".bin")
        return {"canonical_url": url, "response_url": url, "snapshot_path": relative, "sha256": digest,
                "capture_status": "bytes-preserved", "extraction_status": "unsupported-format",
                "size_bytes": len(data), "legal_review_status": "not-reviewed", "current_law_release": False,
                "acquisition_method": "https-original-bytes", "retrieved_at": "2025-01-01T00:00:00Z",
                "references": ["original-reference"], "source_time_version": "not-established",
                "discovery_depth": 1, "discovered_from": [], "links": [], "source_family_ids": [], **changes}

    def audit_row(self, prior):
        return {"canonical_url": prior["canonical_url"], "has_research_text": False,
                "capture_status": "bytes-preserved", "extraction_status": prior["extraction_status"]}

    def recover(self, prior):
        return recovery.recover_one(self.archive, self.output, self.audit_row(prior), prior, {}, time.monotonic() + 45)

    def inputs(self, rows):
        raw = b"".join(recovery.encoded(row) for row in rows)
        manifest = self.archive / "source-originals/manifest.jsonl"
        manifest.parent.mkdir(parents=True, exist_ok=True)
        manifest.write_bytes(raw)
        audit = {"archive_available": True, "inputs": [{"role": "archive", "path": "source-originals/manifest.jsonl",
                                                        "sha256": recovery.sha256(raw)}],
                 "urls": [self.audit_row(row) for row in rows]}
        path = self.archive / "audit.json"
        path.write_bytes(recovery.encoded(audit))
        return manifest, path

    def test_original_identity_formulas_caches_tables_hidden_and_merge_compatibility(self):
        prior = self.prior(workbook_bytes())
        original = self.archive / prior["snapshot_path"]
        before = (original.read_bytes(), original.stat().st_mtime_ns, recovery.encoded(prior))
        row = self.recover(prior)
        self.assertEqual(row["recovery_status"], "captured", row.get("reason"))
        self.assertEqual(row["actual_format"], "xlsx")
        self.assertEqual((original.read_bytes(), original.stat().st_mtime_ns, recovery.encoded(prior)), before)
        self.assertEqual(row["sha256"], prior["sha256"])
        self.assertEqual(row["snapshot_path"], prior["snapshot_path"])
        self.assertEqual(row["preserved_source_provenance"]["source_record"], prior)
        for key in ("canonical_url", "response_url", "acquisition_method", "retrieved_at", "source_time_version", "references"):
            self.assertEqual(row[key], prior[key])
        units = json.loads((self.output / row["text_path"]).read_bytes())
        content = "\n".join(u["text"] for u in units)
        self.assertIn('cached=stored[n]="0.125"', content)
        self.assertIn('cached=stored[n]="0"', content)
        self.assertIn("cached=unavailable (no stored value)", content)
        self.assertIn("=SUM(B3:B4)", content)
        self.assertIn("Hidden source material", content)
        details = row["office_extraction"]
        self.assertEqual(details["sheets"][0]["formula_count"], 3)
        self.assertEqual(details["sheets"][0]["formula_caches_missing"], 1)
        self.assertEqual(details["sheets"][1]["state"], "hidden")
        self.assertEqual(details["sheets"][0]["mergeCell"], [{"ref": "A6:B6"}])
        self.assertEqual(details["tables"][0]["attributes"]["ref"], "A2:B4")
        self.assertTrue(next(u for u in units if u["locator"] == "xlsx:sheet:1:row:4")["row_hidden"])
        self.assertIn("0.0%", details["number_formats"].values())
        self.assertEqual(len(units), len({u["locator"] for u in units}))
        errors = []
        check_artifacts(self.output, row, set(), errors)
        self.assertEqual(errors, [])
        (self.output / recovery.LOG).write_bytes(recovery.encoded(row))
        before_rows, pending, entries, objects = prepare(self.archive, [self.output], [])
        self.assertEqual(len(pending), 1)
        self.assertEqual(len(objects), 2)
        self.assertFalse(row["current_law_release"])
        self.assertEqual(row["legal_review_status"], "not-reviewed")

    def test_empty_workbook_shell_is_failed_without_text(self):
        row = self.recover(self.prior(workbook_bytes(empty=True)))
        self.assertEqual(row["recovery_status"], "failed")
        self.assertEqual(row["reason_code"], "empty-or-insufficient-text")
        self.assertNotIn("text_path", row)
        self.assertTrue(row["original_bytes_unchanged"])

    def test_bad_hash_never_copies_or_claims_text(self):
        prior = self.prior(workbook_bytes())
        (self.archive / prior["snapshot_path"]).write_bytes(b"damaged")
        row = self.recover(prior)
        self.assertEqual(row["reason_code"], "source-hash-mismatch")
        self.assertNotIn("snapshot_path", row)
        self.assertNotIn("text_path", row)

    def test_counterpart_and_other_url_rejected(self):
        prior = self.prior(workbook_bytes(), source_relation={"kind": "public-register-counterpart"})
        self.assertEqual(self.recover(prior)["reason_code"], "non-exact-source")
        audit = self.audit_row(prior)
        audit["canonical_url"] += "?other-version=1"
        with self.assertRaisesRegex(recovery.RecoveryError, "exact audit URL"):
            recovery.source_identity(audit, prior)

    def test_paths_and_overlapping_roots_rejected(self):
        for value in ("../outside", "/outside", "C:outside", "\\\\server\\share", "source-originals/objects/../x", "x:stream", "x./y"):
            with self.subTest(value=value), self.assertRaises(recovery.RecoveryError):
                recovery.safe_path(self.output, value)
        for output in (self.archive, self.archive / "nested", self.archive.parent):
            with self.assertRaises(recovery.RecoveryError):
                recovery.roots(self.archive, output)

    def test_symlink_escape_rejected_when_supported(self):
        self.output.mkdir()
        outside = Path(self.temp.name) / "outside"
        outside.mkdir()
        try:
            (self.output / "source-originals").symlink_to(outside, target_is_directory=True)
        except OSError:
            if sys.platform != "win32":
                self.skipTest("Symlink privilege unavailable")
            created = subprocess.run(["cmd.exe", "/c", "mklink", "/J", str(self.output / "source-originals"), str(outside)],
                                     capture_output=True, timeout=10, creationflags=subprocess.CREATE_NO_WINDOW)
            if created.returncode:
                self.skipTest("Neither symbolic links nor junctions available")
        with self.assertRaises(recovery.RecoveryError):
            recovery.roots(self.archive, self.output)

    def test_selector_uses_latest_preserved_record_and_exact_url(self):
        old = self.prior(workbook_bytes())
        newer = {**old, "extraction_error": "newer record"}
        audit = {"urls": [self.audit_row(old)]}
        selected = recovery.targets_from_report([old, newer], audit)
        self.assertEqual(selected[0][1], newer)
        audit["urls"][0]["canonical_url"] += "?v=2"
        self.assertIsNone(recovery.targets_from_report([old], audit)[0][1])

    def test_stale_audit_fails_before_output(self):
        prior = self.prior(workbook_bytes())
        manifest, audit = self.inputs([prior])
        with manifest.open("ab") as handle:
            handle.write(recovery.encoded(prior))
        with self.assertRaisesRegex(recovery.RecoveryError, "exact source manifest"):
            recovery.run(self.archive, self.output, audit)
        self.assertFalse(self.output.exists())

    def test_append_only_idempotent_run_all_candidates_have_outcomes(self):
        first = self.prior(workbook_bytes())
        blank = self.prior(workbook_bytes(empty=True), url="https://example.gov.au/empty.xlsx")
        _, audit = self.inputs([first, blank])
        result = recovery.run(self.archive, self.output, audit)
        self.assertEqual((result["candidate_count"], result["recovered_count"], result["failed_count"]), (2, 1, 1))
        log = self.output / recovery.LOG
        before = log.read_bytes()
        again = recovery.run(self.archive, self.output, audit)
        self.assertEqual(again["appended_this_run"], 0)
        self.assertEqual(log.read_bytes(), before)
        self.assertEqual(len(read_jsonl(log)), 2)

    def test_output_corruption_is_detected_on_replay(self):
        prior = self.prior(workbook_bytes())
        _, audit = self.inputs([prior])
        recovery.run(self.archive, self.output, audit)
        row = read_jsonl(self.output / recovery.LOG)[0]
        (self.output / row["text_path"]).write_bytes(b"[]")
        with self.assertRaisesRegex(recovery.RecoveryError, "Output hash mismatch"):
            recovery.run(self.archive, self.output, audit)

    def test_replay_cannot_substitute_a_different_source_record(self):
        prior = self.prior(workbook_bytes())
        _, audit = self.inputs([prior])
        recovery.run(self.archive, self.output, audit)
        log = self.output / recovery.LOG
        row = read_jsonl(log)[0]
        row["preserved_source_provenance"]["source_record"]["canonical_url"] += "?different=1"
        log.write_bytes(recovery.encoded(row))
        with self.assertRaisesRegex(recovery.RecoveryError, "selected source record"):
            recovery.run(self.archive, self.output, audit)

    def test_time_budget_remains_failed_and_preserves_original(self):
        prior = self.prior(workbook_bytes())
        row = recovery.recover_one(self.archive, self.output, self.audit_row(prior), prior, {}, 0)
        self.assertEqual(row["recovery_status"], "failed")
        self.assertEqual(row["reason_code"], "time-limit")
        self.assertNotIn("text_path", row)
        self.assertTrue(row["original_bytes_unchanged"])

    def test_zip_limits_and_unsafe_members(self):
        data = workbook_bytes()
        with patch.object(recovery, "MAX_ZIP_BYTES", 1), self.assertRaises(recovery.RecoveryError):
            recovery.checked_package(data)
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, "w") as package:
            package.writestr("../escape", "bad")
        with self.assertRaises(recovery.RecoveryError):
            recovery.checked_package(raw.getvalue())

    def test_unknown_ole_and_non_office_shell_rejected(self):
        for data in (bytes.fromhex("d0cf11e0a1b11ae1") + b"empty", b"<html>unavailable</html>"):
            with self.assertRaises(recovery.RecoveryError):
                recovery.detect_format(data)

    def test_rtf_windows_runtime_preserves_unicode_and_table_cells(self):
        if sys.platform != "win32":
            self.skipTest("Windows WPF test")
        data = (r"{\rtf1\ansi\ansicpg1252\uc1 " + TEXT + r"\par Unicode \u945? and escaped \{braces\}.\par "
                r"\trowd\cellx1800\cellx3600\intbl Duty\cell Evidence\cell\row}").encode("ascii")
        row = self.recover(self.prior(data, url="https://example.gov.au/source.rtf"))
        self.assertEqual(row["recovery_status"], "captured", row.get("reason"))
        units = json.loads((self.output / row["text_path"]).read_bytes())
        self.assertIn(chr(945), "\n".join(u["text"] for u in units))
        self.assertIn("{braces}", "\n".join(u["text"] for u in units))
        self.assertTrue(any(":row:1:cell:1" in u["locator"] and u["text"] == "Duty" for u in units))


if __name__ == "__main__":
    unittest.main()
