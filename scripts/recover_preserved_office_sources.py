"""Recover exact preserved Office sources offline into a disjoint evidence batch.

Uses openpyxl for XLSX values/formulas, lxml only for stored caches and package
metadata, and Windows WPF/IFilter for RTF/DOC. No recalculation or legal review.
"""

from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import io
import json
import os
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

sys.dont_write_bytecode = True
from lxml import etree
import openpyxl
from collect_source_originals import now, read_jsonl, save_object
from source_manifest import select_manifest_rows
from validate_source_originals import ARTIFACTS


REVISION = 1
METHOD = "preserved-office-offline-v1"
LOG = "source-originals/recovery-office.jsonl"
STATUSES = {"unsupported-format", "extraction-error"}
STAGE = Path(__file__).resolve().parents[1]
HELPER = STAGE / "scripts/read_legacy_office.ps1"
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MAX_ZIP_BYTES = 500_000_000
MAX_TEXT_BYTES = 500_000_000
MAX_CELLS = 8_000_000
BOUNDARIES = [
    "Recovered text is research evidence only; conversion is not legal verification.",
    "Exact preserved bytes establish source identity, not historical or current legal applicability.",
    "Extraction completeness and visual layout have not been independently reviewed.",
]


class RecoveryError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def digest_file(path):
    with Path(path).open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def encoded(value):
    return (json.dumps(value, ensure_ascii=True, sort_keys=True) + "\n").encode("utf-8")


def safe_path(root, relative, objects=False):
    value = str(relative)
    parts = value.replace("\\", "/").split("/")
    if (not value or PureWindowsPath(value).drive or value.startswith(("/", "\\"))
            or any(p in {"", ".", ".."} or ":" in p or p.endswith((".", " ")) for p in parts)):
        raise RecoveryError("unsafe-path", "Unsafe relative path: " + value)
    root = Path(root).resolve()
    base = (root / "source-originals/objects").resolve() if objects else root
    path = root.joinpath(*parts).resolve()
    if not base.is_relative_to(root) or not path.is_relative_to(base) or path == base:
        raise RecoveryError("unsafe-path", "Resolved path escapes its permitted root")
    return path


def roots(archive_root, output_root):
    archive_root, output_root = Path(archive_root).resolve(), Path(output_root).resolve()
    if archive_root.is_relative_to(output_root) or output_root.is_relative_to(archive_root):
        raise RecoveryError("overlapping-roots", "Output and archive roots must be disjoint")
    safe_path(archive_root, "source-originals/objects/check", objects=True)
    safe_path(output_root, "source-originals/objects/check", objects=True)
    safe_path(output_root, LOG)
    return archive_root, output_root


def preserve(root, data, suffix):
    if not re.fullmatch(r"\.[a-z0-9]+", suffix):
        raise RecoveryError("unsafe-suffix", "Invalid object suffix")
    digest = sha256(data)
    relative = f"source-originals/objects/{digest[:2]}/{digest}{suffix}"
    safe_path(root, relative, objects=True)
    return save_object(root, data, suffix)


def verified_original(archive_root, prior):
    path = safe_path(archive_root, prior.get("snapshot_path", ""), objects=True)
    digest = prior.get("sha256", "")
    if not re.fullmatch("[0-9a-f]{64}", digest):
        raise RecoveryError("invalid-source-hash", "Original needs a lowercase SHA-256")
    if path.name != digest + path.suffix or path.parent.name != digest[:2]:
        raise RecoveryError("invalid-hash-path", "Original path must be content addressed")
    data = path.read_bytes()
    if sha256(data) != digest:
        raise RecoveryError("source-hash-mismatch", "Preserved bytes do not match the source record")
    if prior.get("size_bytes", len(data)) != len(data):
        raise RecoveryError("source-size-mismatch", "Preserved size does not match the source record")
    return data, path


def xml(data):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    tree = etree.fromstring(data, parser)
    if tree.getroottree().docinfo.doctype:
        raise RecoveryError("xml-dtd", "DTD-bearing Office XML is not accepted")
    return tree


def checked_package(data):
    package = zipfile.ZipFile(io.BytesIO(data))
    infos = package.infolist()
    names = [i.filename for i in infos]
    if len(names) != len(set(names)) or len(names) > 20000:
        package.close()
        raise RecoveryError("zip-members", "Duplicate or excessive ZIP members")
    if sum(i.file_size for i in infos) > MAX_ZIP_BYTES:
        package.close()
        raise RecoveryError("zip-size-limit", "Office package exceeds bounded 500 MB expansion")
    for item in infos:
        parts = item.filename.rstrip("/").split("/")
        if (item.flag_bits & 1 or item.filename.startswith("/") or "\\" in item.filename
                or any(p in {"", ".", ".."} or ":" in p for p in parts)):
            package.close()
            raise RecoveryError("zip-path", "Unsafe/encrypted Office ZIP member")
    return package


def detect_format(data):
    if data.lstrip().startswith(b"{\\rtf"):
        return "rtf"
    if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        # OLE signature alone does not prove this is a Word document.
        if "WordDocument".encode("utf-16-le") in data:
            return "doc"
        raise RecoveryError("unsupported-ole", "OLE object has no WordDocument stream name")
    if data.startswith(b"PK"):
        with checked_package(data) as package:
            if "xl/workbook.xml" in package.namelist():
                return "xlsx"
        raise RecoveryError("unsupported-zip", "ZIP is not a supported workbook")
    raise RecoveryError("unsupported-bytes", "Bytes are not RTF, legacy Word, or XLSX")


def check_deadline(deadline):
    if time.monotonic() > deadline:
        raise RecoveryError("time-limit", "Bounded extraction time exhausted")


def literal(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return json.dumps(value, ensure_ascii=True, separators=(",", ":"), allow_nan=False)


class UnitWriter:
    """Stream row units so large workbooks do not accumulate millions of cells."""
    def __init__(self, path):
        self.handle = path.open("xb")
        self.handle.write(b"[")
        self.count = self.characters = self.substantive = self.size = 0

    def add(self, unit, source_content=True):
        if not unit["text"].strip():
            return
        raw = encoded(unit).rstrip(b"\n")
        self.size += len(raw) + 2
        if self.size > MAX_TEXT_BYTES:
            raise RecoveryError("text-size-limit", "Extracted text exceeds 500 MB limit")
        self.handle.write((b",\n" if self.count else b"\n") + raw)
        self.count += 1
        self.characters += len(unit["text"])
        if source_content:
            self.substantive += sum(c.isalnum() for c in unit["text"])

    def close(self):
        self.handle.write(b"\n]\n")
        self.handle.close()


def worksheet_rows(package, part, metadata, deadline):
    """Read stored formula caches and structural metadata; never evaluate formulas."""
    with package.open(part) as stream:
        context = etree.iterparse(stream, events=("end",), resolve_entities=False,
                                  no_network=True, load_dtd=False)
        for _, node in context:
            if node.tag == S + "row":
                check_deadline(deadline)
                formulas = {}
                for cell in node.findall(S + "c"):
                    metadata["stored_cell_count"] += 1
                    if metadata["stored_cell_count"] > MAX_CELLS:
                        raise RecoveryError("cell-limit", "Worksheet exceeds 8 million stored cells")
                    formula = cell.find(S + "f")
                    if formula is not None:
                        value = cell.find(S + "v")
                        raw = value.text if value is not None else None
                        formulas[cell.get("r")] = (cell.get("t", "n"), raw)
                        metadata["formula_count"] += 1
                        metadata["formula_caches_missing"] += raw is None
                number = int(node.get("r"))
                hidden = node.get("hidden") == "1"
                node.clear()
                while node.getprevious() is not None:
                    del node.getparent()[0]
                yield number, formulas, hidden
            elif node.tag in {S + "mergeCell", S + "col", S + "tablePart", S + "hyperlink"}:
                metadata.setdefault(etree.QName(node).localname, []).append(dict(node.attrib))
            elif node.tag in {S + "oddHeader", S + "evenHeader", S + "firstHeader",
                              S + "oddFooter", S + "evenFooter", S + "firstFooter"}:
                metadata.setdefault("header_footer_codes", {})[etree.QName(node).localname] = node.text
            elif node.tag == S + "dimension":
                metadata["declared_dimension"] = node.get("ref")
        if context.root.getroottree().docinfo.doctype:
            raise RecoveryError("xml-dtd", "DTD-bearing worksheet is not accepted")


def workbook_parts(package):
    workbook = xml(package.read("xl/workbook.xml"))
    relationships = xml(package.read("xl/_rels/workbook.xml.rels"))
    targets = {}
    for rel in relationships:
        if rel.get("Type", "").endswith("/worksheet"):
            target = rel.get("Target", "")
            if rel.get("TargetMode") == "External" or "\\" in target or ":" in target or ".." in target.split("/"):
                raise RecoveryError("unsafe-relationship", "Worksheet target is not an internal package part")
            part = target.lstrip("/") if target.startswith("/") else "xl/" + target
            if part not in package.namelist():
                raise RecoveryError("missing-worksheet", "Referenced worksheet part is missing")
            targets[rel.get("Id")] = part
    sheets = []
    for sheet in workbook.findall(S + "sheets/" + S + "sheet"):
        if sheet.get(R + "id") not in targets:
            raise RecoveryError("unsupported-sheet", "Non-worksheet sheet needs a separate parser")
        sheets.append({"name": sheet.get("name"), "state": sheet.get("state", "visible"),
                       "part": targets[sheet.get(R + "id")]})
    defined = [{"attributes": dict(n.attrib), "text": n.text} for n in workbook.findall(S + "definedNames/" + S + "definedName")]
    return sheets, defined


def extract_xlsx(data, writer, deadline):
    with checked_package(data) as package:
        sheets, defined = workbook_parts(package)
        metadata = {"parser": "openpyxl " + openpyxl.__version__, "sheets": [], "number_formats": {},
                    "defined_names": defined, "formula_policy": "No recalculation. Stored cached values may be stale. Missing caches stay unavailable.",
                    "limits": ["Cell rows, formulas, stored caches, sheet visibility, merge ranges and table definitions are recovered.",
                               "Blank/styled-only cells are omitted, not converted to zero; coordinates preserve gaps.",
                               "Numeric cached tokens are as stored in XML; number-format codes define their display units.",
                               "Dates/values use openpyxl interpretation; displayed formatting and Excel behavior are unverified.",
                               "Drawings, images, charts, embedded objects, comments and external-link contents are not text-extracted.",
                               "Shared formulas are expanded by openpyxl; original formula representation remains in preserved bytes."],
                    "unextracted_parts": [p for p in package.namelist() if re.search(r"/(drawings|charts|media|embeddings|externalLinks)/|comments\d*\.xml$", p)]}
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if workbook.sheetnames != [s["name"] for s in sheets]:
                raise RecoveryError("sheet-identity", "Workbook parser and package sheet identities disagree")
            formats = {}
            for sheet_index, source in enumerate(sheets, 1):
                check_deadline(deadline)
                sheet = workbook[source["name"]]
                sheet.reset_dimensions()
                info = {**source, "stored_cell_count": 0, "formula_count": 0, "formula_caches_missing": 0,
                        "recovered_cell_count": 0, "recovered_row_count": 0}
                raw_rows = iter(worksheet_rows(package, source["part"], info, deadline))
                raw = next(raw_rows, None)
                for row_number, cells in enumerate(sheet.iter_rows(), 1):
                    check_deadline(deadline)
                    if raw is not None and raw[0] < row_number:
                        raise RecoveryError("row-identity", "Worksheet parser skipped a stored row")
                    cached, hidden = (raw[1], raw[2]) if raw is not None and raw[0] == row_number else ({}, False)
                    values = []
                    seen_formulas = set()
                    for cell in cells:
                        value = cell.value
                        if value is None:
                            continue
                        code = cell.number_format
                        if code not in formats:
                            formats[code] = len(formats) + 1
                        fmt = formats[code]
                        if cell.data_type == "f":
                            if cell.coordinate not in cached:
                                raise RecoveryError("formula-identity", "Formula does not match its source cell")
                            seen_formulas.add(cell.coordinate)
                            cache_type, cache = cached[cell.coordinate]
                            if not isinstance(value, str):
                                value = {"kind": type(value).__name__, **dict(value), "text": getattr(value, "text", None)}
                            result = "unavailable (no stored value)" if cache is None else f"stored[{cache_type}]={literal(cache)}"
                            values.append(f"{cell.coordinate} [formula;format:{fmt}]: {literal(value)}; cached={result}")
                        else:
                            values.append(f"{cell.coordinate} [{cell.data_type};format:{fmt}]: {literal(value)}")
                    if seen_formulas != set(cached):
                        raise RecoveryError("formula-identity", "Stored formulas were omitted by the workbook parser")
                    if values:
                        writer.add({"locator": f"xlsx:sheet:{sheet_index}:row:{row_number}",
                                    "sheet": source["name"], "source_part": source["part"],
                                    "row_hidden": hidden, "text": "\n".join(values)})
                        info["recovered_cell_count"] += len(values)
                        info["recovered_row_count"] += 1
                    if raw is not None and raw[0] == row_number:
                        raw = next(raw_rows, None)
                if raw is not None:
                    raise RecoveryError("row-identity", "Stored rows remain unread")
                metadata["sheets"].append(info)
            metadata["number_formats"] = {str(number): code for code, number in formats.items()}
            metadata["tables"] = []
            for part in package.namelist():
                if re.fullmatch(r"xl/tables/table\d+\.xml", part):
                    table = xml(package.read(part))
                    metadata["tables"].append({"part": part, "attributes": dict(table.attrib),
                                               "xml": etree.tostring(table, encoding="unicode")})
            writer.add({"locator": "xlsx:workbook:structure", "text": json.dumps(metadata, ensure_ascii=True, sort_keys=True)}, source_content=False)
            metadata["all_sheets_read"] = True
            return metadata
        finally:
            workbook.close()


def extract_legacy(copy_path, kind, writer, deadline, scratch):
    executable = shutil.which("pwsh") or shutil.which("powershell")
    if not executable or not HELPER.is_file():
        raise RecoveryError("legacy-parser-unavailable", "Windows parser helper/runtime is unavailable")
    # IFilter dispatches on extension; only this disposable copy receives .doc.
    typed_path = Path(scratch) / ("original." + kind)
    shutil.copyfile(copy_path, typed_path)
    try:
        command = [executable, "-NoProfile", "-NonInteractive", "-STA", "-File", str(HELPER),
                   "-Source", str(typed_path), "-Format", kind]
        completed = subprocess.run(command, capture_output=True, timeout=max(1, min(45, deadline - time.monotonic())),
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        if completed.returncode:
            raise RecoveryError("legacy-parser-error", completed.stderr.decode("utf-8", errors="replace")[-3000:])
        parsed = json.loads(completed.stdout.decode("utf-8-sig"))
        for unit in parsed["units"]:
            writer.add(unit)
        limits = (["WPF imports the RTF body and table cell text, including row/column spans.",
                   "Headers, footers, footnotes, field instructions, images, revision marks and pagination may be omitted or changed by WPF.",
                   "List numbering and nested table geometry are not independently verified."] if kind == "rtf" else
                  ["Windows IFilter emits source chunks and text separators; these are not page or legal-section locators.",
                   "Table text is retained in parser order; full cell geometry, formatting, headers, notes and embedded objects are not established."])
        return {"parser": parsed["parser"], "limits": limits, "helper_sha256": digest_file(HELPER),
                "original_copy_read_only": True, "office_application_started": False}
    finally:
        typed_path.unlink(missing_ok=True)


def targets_from_report(manifest_rows, report):
    latest, preserved = select_manifest_rows(manifest_rows)
    targets, seen = [], set()
    for audit in report["urls"]:
        if (audit.get("has_research_text") is not False or audit.get("capture_status") != "bytes-preserved"
                or audit.get("extraction_status") not in STATUSES):
            continue
        url = audit["canonical_url"]
        if url in seen:
            raise RecoveryError("duplicate-audit-url", "Audit repeats a candidate identity")
        seen.add(url)
        prior = preserved.get(url) or latest.get(url)
        targets.append((audit, deepcopy(prior)))
    return targets


def source_identity(audit, prior):
    url = audit["canonical_url"]
    if not prior or prior.get("canonical_url") != url:
        raise RecoveryError("missing-exact-source", "No record for the exact audit URL")
    if prior.get("capture_status") != "bytes-preserved" or prior.get("extraction_status") not in STATUSES:
        raise RecoveryError("source-status-mismatch", "Selected preserved record does not match the audited missing-text status")
    if prior.get("source_relation") or prior.get("recovery_for_url") not in (None, url):
        raise RecoveryError("non-exact-source", "Alternative or counterpart evidence is not an exact-source recovery")


def recover_one(archive_root, output_root, audit, prior, provenance, deadline):
    archive_root, output_root = roots(archive_root, output_root)
    row = deepcopy(prior or {"canonical_url": audit["canonical_url"]})
    for path_key, hash_key in ARTIFACTS:
        row.pop(path_key, None)
        row.pop(hash_key, None)
    for key in ("extraction_error", "recovery_import_id", "recovery_imported_at", "invalidates_prior_text"):
        row.pop(key, None)
    row.update(recovery_status="failed", capture_status="failed", extraction_status="extraction-error",
               extraction_method=METHOD, extraction_revision=REVISION, recovered_at=now(),
               text_character_count=0, legal_review_status="not-reviewed", current_law_release=False,
               extraction_completeness="not-established", original_bytes_unchanged=False,
               preserved_source_provenance={**provenance, "source_record": deepcopy(prior), "audit_record": deepcopy(audit)})
    original = None
    try:
        source_identity(audit, prior)
        data, original = verified_original(archive_root, prior)
        before_stat = original.stat()
        snapshot, digest = preserve(output_root, data, original.suffix)
        row.update(snapshot_path=snapshot, sha256=digest, capture_status="bytes-preserved", original_copied=True)
        kind = detect_format(data)
        row["actual_format"] = kind
        check_deadline(deadline)
        temporary = safe_path(output_root, "work")
        temporary.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="office-", dir=temporary) as scratch:
            text_file = Path(scratch) / "units.json"
            writer = UnitWriter(text_file)
            try:
                if kind == "xlsx":
                    extraction = extract_xlsx(data, writer, deadline)
                else:
                    extraction = extract_legacy(safe_path(output_root, snapshot, objects=True), kind, writer, deadline, scratch)
            finally:
                writer.close()
            row["office_extraction"] = extraction
            if writer.substantive < 80:
                raise RecoveryError("empty-or-insufficient-text", "Below 80 substantive source characters; metadata alone is not a recovery")
            # Source bytes, size and modification time must survive extraction unchanged.
            after_stat = original.stat()
            if digest_file(original) != prior["sha256"] or (before_stat.st_size, before_stat.st_mtime_ns) != (after_stat.st_size, after_stat.st_mtime_ns):
                raise RecoveryError("source-changed", "Original changed during extraction")
            text_path, text_hash = preserve(output_root, text_file.read_bytes(), ".json")
            row.update(recovery_status="captured", extraction_status="extracted-unreviewed",
                       text_path=text_path, text_sha256=text_hash, text_character_count=writer.characters,
                       recovered_unit_count=writer.count, substantive_character_count=writer.substantive,
                       content_class=kind + "-preserved-office-text-unreviewed",
                       extraction_completeness="completeness-unreviewed", original_bytes_unchanged=True)
    except Exception as exc:
        row.update(reason_code=getattr(exc, "code", "parser-error"), reason=f"{type(exc).__name__}: {exc}")
        if original is not None:
            row["original_bytes_unchanged"] = digest_file(original) == prior["sha256"]
    return row


def recovery_key(audit, prior):
    return sha256(encoded({"url": audit["canonical_url"], "source": prior, "method": METHOD, "revision": REVISION}))


def verify_outcome(archive_root, output_root, row):
    prior = row["preserved_source_provenance"]["source_record"]
    audit = row["preserved_source_provenance"]["audit_record"]
    if row.get("office_recovery_key") is not None and row["office_recovery_key"] != recovery_key(audit, prior):
        raise RecoveryError("replay-identity", "Saved outcome key does not match its exact source record")
    if row.get("snapshot_path"):
        if not prior or row["canonical_url"] != prior["canonical_url"] or row["sha256"] != prior["sha256"]:
            raise RecoveryError("output-identity", "Output does not preserve exact source identity")
        if digest_file(safe_path(archive_root, prior["snapshot_path"], objects=True)) != prior["sha256"]:
            raise RecoveryError("source-changed", "Original no longer matches its preserved hash")
    for path_key, hash_key in ARTIFACTS:
        if path_key in row:
            path = safe_path(output_root, row[path_key], objects=True)
            if digest_file(path) != row.get(hash_key):
                raise RecoveryError("output-integrity", "Output hash mismatch: " + row[path_key])
    if row["recovery_status"] == "captured" and not row.get("text_path"):
        raise RecoveryError("missing-output-text", "Captured outcome has no text")
    if row["recovery_status"] != "captured" and row.get("text_path"):
        raise RecoveryError("failed-has-text", "Failed outcome must not claim recovered text")
    if row.get("legal_review_status") != "not-reviewed" or row.get("current_law_release") is not False:
        raise RecoveryError("release-state", "Recovery cannot promote legal review state")


def run(archive_root, output_root, audit_path, per_source_seconds=180, total_seconds=1500):
    archive_root, output_root = roots(archive_root, output_root)
    audit_path = Path(audit_path).resolve()
    audit_raw = audit_path.read_bytes()
    report = json.loads(audit_raw)
    manifest = safe_path(archive_root, "source-originals/manifest.jsonl")
    manifest_raw = manifest.read_bytes()
    manifest_hash = sha256(manifest_raw)
    audit_manifest = [r for r in report.get("inputs", []) if r.get("role") == "archive" and r.get("path") == "source-originals/manifest.jsonl"]
    if report.get("archive_available") is not True or len(audit_manifest) != 1 or audit_manifest[0].get("sha256") != manifest_hash:
        raise RecoveryError("audit-manifest-mismatch", "Audit must identify the exact source manifest by hash")
    candidates = targets_from_report([json.loads(line) for line in manifest_raw.decode("utf-8-sig").splitlines() if line.strip()], report)
    provenance = {"archive_root": str(archive_root), "manifest_path": "source-originals/manifest.jsonl",
                  "manifest_sha256": manifest_hash, "audit_path": str(audit_path), "audit_sha256": sha256(audit_raw)}
    log = safe_path(output_root, LOG)
    existing = {r.get("office_recovery_key"): r for r in read_jsonl(log)}
    log.parent.mkdir(parents=True, exist_ok=True)
    total_deadline = time.monotonic() + total_seconds
    results, appended = [], 0
    for index, (audit, prior) in enumerate(candidates, 1):
        key = recovery_key(audit, prior)
        if key in existing:
            outcome = existing[key]
            if (outcome["preserved_source_provenance"]["source_record"] != prior
                    or outcome["canonical_url"] != audit["canonical_url"]):
                raise RecoveryError("replay-identity", "Saved outcome does not match the selected source record")
            verify_outcome(archive_root, output_root, outcome)
        else:
            deadline = min(total_deadline, time.monotonic() + per_source_seconds)
            outcome = recover_one(archive_root, output_root, audit, prior, provenance, deadline)
            outcome["office_recovery_key"] = key
            verify_outcome(archive_root, output_root, outcome)
            with log.open("ab") as handle:
                handle.write(encoded(outcome))
            appended += 1
        results.append(outcome)
        print(json.dumps({"candidate": index, "of": len(candidates), "url": audit["canonical_url"],
                          "format": outcome.get("actual_format"), "status": outcome["recovery_status"],
                          "reason": outcome.get("reason_code")}), flush=True)
    if digest_file(manifest) != manifest_hash or digest_file(audit_path) != provenance["audit_sha256"]:
        raise RecoveryError("input-changed", "Audit or canonical manifest changed during recovery")
    for row in results:
        verify_outcome(archive_root, output_root, row)
    return {"generated_at": now(), "method": METHOD, "revision": REVISION, **provenance,
            "output_root": str(output_root), "candidate_count": len(candidates),
            "recovered_count": sum(r["recovery_status"] == "captured" for r in results),
            "failed_count": sum(r["recovery_status"] != "captured" for r in results),
            "audit_status_counts": dict(Counter(a["extraction_status"] for a, _ in candidates)),
            "format_counts": dict(Counter(r.get("actual_format", "unidentified") for r in results)),
            "appended_this_run": appended, "recovery_log": LOG,
            "source_manifest_unchanged": True, "audit_unchanged": True,
            "copied_original_count": sum(bool(r.get("original_copied")) for r in results),
            "originals_hash_verified_unchanged": sum(bool(r.get("original_bytes_unchanged")) for r in results),
            "output_hash_verification": "passed", "network_requests": 0,
            "legal_review_status": "not-reviewed", "current_law_release": False,
            "limits": {"per_source_seconds": per_source_seconds, "total_seconds": total_seconds,
                       "max_zip_expanded_bytes": MAX_ZIP_BYTES, "max_text_bytes": MAX_TEXT_BYTES},
            "boundaries": BOUNDARIES, "outcomes": results}


def write_report(report, prefix):
    prefix = Path(prefix).absolute()
    archive = Path(report["archive_root"]).resolve()
    for extension in (".json", ".md"):
        path = prefix.with_suffix(extension).resolve()
        if path.is_relative_to(archive):
            raise RecoveryError("report-in-archive", "Report cannot be written inside the archive")
    prefix.parent.mkdir(parents=True, exist_ok=True)
    prefix.with_suffix(".json").write_bytes((json.dumps(report, ensure_ascii=True, indent=2) + "\n").encode())
    lines = ["# Preserved Office Source Recovery", "",
             f"Recovered: {report['recovered_count']}. Failed: {report['failed_count']}. Candidates: {report['candidate_count']}.", "",
             "Canonical archive is read only. No network acquisition, recalculation, or legal classification changes.",
             "All outcomes remain not-reviewed; current_law_release is false.", "",
             f"Manifest SHA-256: `{report['manifest_sha256']}`.",
             f"Originals hash-verified unchanged: {report['originals_hash_verified_unchanged']}. Output-object verification passed.", "",
             "## Outcomes", "", "| Source URL | Format | Outcome | Units | Limitation / failure |",
             "| --- | --- | --- | ---: | --- |"]
    for row in report["outcomes"]:
        detail = row.get("reason") or " ".join(row.get("office_extraction", {}).get("limits", []))
        lines.append(f"| {row['canonical_url']} | {row.get('actual_format', 'unknown')} | {row['recovery_status']} | {row.get('recovered_unit_count', 0)} | {detail.replace('|', '/').replace(chr(10), ' ')} |")
    lines += ["", "## Boundaries", ""] + ["- " + item for item in BOUNDARIES]
    lines += ["", "The JSON report and append-only log retain exact original source records, hashes, locators, parser details and per-workbook counts.", ""]
    prefix.with_suffix(".md").write_text("\n".join(lines), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True, help="Disjoint evidence-batch root")
    parser.add_argument("--audit", type=Path, default=STAGE / "review/results/delivery-readiness.json")
    parser.add_argument("--report-prefix", type=Path)
    parser.add_argument("--per-source-seconds", type=float, default=180)
    parser.add_argument("--total-seconds", type=float, default=1500)
    args = parser.parse_args()
    report = run(args.archive_root, args.output_root, args.audit, args.per_source_seconds, args.total_seconds)
    write_report(report, args.report_prefix or args.output_root / "summary")
    print(json.dumps({k: report[k] for k in ("candidate_count", "recovered_count", "failed_count", "appended_this_run")}))


if __name__ == "__main__":
    main()
