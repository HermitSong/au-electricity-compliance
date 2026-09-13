#!/usr/bin/env python3
"""Preserve public source originals without promoting them to verified law."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict, deque
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
import sys
import threading
import time
import zipfile
from urllib.error import HTTPError
from urllib.parse import quote, urldefrag, urljoin, urlparse, urlunparse
from urllib.request import HTTPRedirectHandler, Request, build_opener
from urllib.robotparser import RobotFileParser

from lxml import etree, html
from pypdf import PdfReader
from source_manifest import select_manifest_rows


USER_AGENT = "AU-Power-Compliance-KB/1.0 public-source-preservation"
URL_PATTERN = re.compile(r"https?://[^\s<>\"`|]+")
ATTACHMENT = re.compile(r"\.(?:pdf|docx?|xlsx?|csv|rtf|txt|xml)(?:$|[?.])|/view/pdf/", re.I)
NON_DOCUMENT = re.compile(r"privacy|cookie|accessibility|terms.of.use|annual.report", re.I)
CHALLENGE_TITLE = re.compile(r"access denied|forbidden|just a moment|verify you are|page not found|^404\b|attention required|request rejected|^sign in\b|^log in\b", re.I)
EXTRA_AUTHORITY_HOSTS = {
    "aemo.com.au", "sapowernetworks.com.au", "ausgrid.com.au", "ausnetservices.com.au",
    "energex.com.au", "powercor.com.au", "media.powercor.com.au", "endeavourenergy.com.au",
    "essentialenergy.com.au", "tasnetworks.com.au", "transgrid.com.au",
}
OBJECT_LOCKS = defaultdict(threading.Lock)


def now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]


def normalise_url(value: str) -> str:
    value = urldefrag(value.strip())[0]
    parsed = urlparse(value)
    return urlunparse(parsed._replace(scheme=parsed.scheme.lower(), netloc=parsed.netloc.lower()))


def host_key(url: str) -> str:
    return (urlparse(url).hostname or "").lower().removeprefix("www.")


def transport_url(url: str) -> str:
    parsed = urlparse(url)
    return urlunparse(parsed._replace(
        path=quote(parsed.path, safe="/%:@!$&'()*+,;=-._~"),
        query=quote(parsed.query, safe="/?%:@!$&'()*+,;=-._~")))


def allowed_url(url: str, hosts: set[str]) -> bool:
    try:
        parsed = urlparse(url)
        return (parsed.scheme == "https" and not parsed.username and not parsed.password
                and parsed.port in (None, 443) and bool(parsed.hostname)
                and (host_key(url).endswith(".gov.au") or host_key(url) in hosts))
    except ValueError:
        return False


def seeds(root: Path) -> tuple[dict[str, dict], set[str]]:
    sys.path.insert(0, str(root / "scripts"))
    from build_source_artifact_ledger import collect_canonical_urls
    references, families = collect_canonical_urls(root)
    hosts = {host_key(url) for url in references} | EXTRA_AUTHORITY_HOSTS
    entries: dict[str, dict] = {}

    def add(url: str, reference: str, family_ids=()) -> None:
        url = normalise_url(url)
        item = entries.setdefault(url, {"canonical_url": url, "references": [], "source_family_ids": [],
                                        "discovery_depth": 0, "discovered_from": []})
        item["references"] = sorted(set(item["references"]) | {reference})
        item["source_family_ids"] = sorted(set(item["source_family_ids"]) | set(family_ids))

    for url, refs in references.items():
        for ref in refs:
            add(url, ref, families.get(url, set()))
    paths = sorted((root / "knowledge-base").rglob("*.md")) + [root / "official-documents" / "SOURCES.md"]
    for path in paths:
        for match in URL_PATTERN.finditer(path.read_text(encoding="utf-8-sig")):
            url = match.group().rstrip(".,;:")
            while url.endswith(")") and url.count(")") > url.count("("):
                url = url[:-1]
            url = url.rstrip("]")
            add(url, f"authored-source:{path.relative_to(root).as_posix()}")
    return entries, hosts


def save_object(root: Path, data: bytes, suffix: str) -> tuple[str, str]:
    digest = hashlib.sha256(data).hexdigest()
    relative = Path("source-originals") / "objects" / digest[:2] / (digest + suffix)
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    with OBJECT_LOCKS[str(path)]:
        if path.exists():
            if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                raise ValueError("Existing content-addressed object has a checksum mismatch")
        else:
            with path.open("xb") as handle:
                handle.write(data)
    return relative.as_posix(), digest


def extract(data: bytes, content_type: str, url: str) -> dict:
    result = {"title": "", "units": [], "links": [], "extraction_status": "unsupported-format",
              "page_count": None, "pages_without_text": [], "content_class": "unclassified"}
    if data.startswith(b"%PDF-"):
        reader = PdfReader(io.BytesIO(data))
        result["title"] = str((reader.metadata or {}).get("/Title") or "")
        result["page_count"] = len(reader.pages)
        for number, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").replace("\x00", "").strip()
            if text:
                result["units"].append({"locator": f"page:{number}", "text": text})
            else:
                result["pages_without_text"].append(number)
        result["extraction_status"] = "needs-ocr-or-blank-page-review" if result["pages_without_text"] else "extracted-unreviewed"
        result["content_class"] = "pdf-document"
        return result
    if data.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            if sum(info.file_size for info in package.infolist()) > 100_000_000:
                raise ValueError("Office package exceeds uncompressed extraction limit")
            names = package.namelist()
            if "xl/workbook.xml" in names:
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                try:
                    for sheet in workbook:
                        for number, row in enumerate(sheet.iter_rows(), 1):
                            values = [f"{cell.coordinate}: {cell.value}" for cell in row if cell.value is not None]
                            if values:
                                result["units"].append({"locator": f"xlsx:{sheet.title}:row:{number}", "text": "\n".join(values)})
                finally:
                    workbook.close()
                result["extraction_status"] = "extracted-unreviewed" if result["units"] else "empty-document"
                result["content_class"] = "xlsx-cell-values-and-unevaluated-formulas-unreviewed"
                return result
            if "word/document.xml" not in names:
                return result
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            parts = [name for name in names if re.fullmatch(r"word/(document|footnotes|endnotes|header\d+|footer\d+)\.xml", name)]
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for part in sorted(parts):
                tree = etree.fromstring(package.read(part), parser=parser)
                for number, paragraph in enumerate(tree.xpath("//w:p", namespaces=ns), 1):
                    text = "".join(paragraph.xpath(".//w:t/text()", namespaces=ns)).strip()
                    if text:
                        result["units"].append({"locator": f"docx:{part}:paragraph:{number}", "text": text})
            result["extraction_status"] = "extracted-unreviewed" if result["units"] else "empty-document"
            result["content_class"] = "docx-document-layout-unreviewed"
            return result
    if "html" in content_type or data.lstrip().lower().startswith((b"<!doctype html", b"<html")):
        tree = html.fromstring(data)
        result["title"] = " ".join(tree.xpath("//title/text()")).strip()
        for anchor in tree.xpath("//a[@href]"):
            href = urljoin(url, anchor.get("href"))
            result["links"].append({"url": normalise_url(href), "label": " ".join(anchor.itertext()).strip()})
        if CHALLENGE_TITLE.search(result["title"]):
            result["extraction_status"] = "blocked-or-error-page"
            return result
        mains = tree.xpath("//main | //article | //*[@id='main-layout'] | //*[@role='main']")
        body = max(mains, key=lambda node: len(" ".join(node.itertext()))) if mains else tree
        for node in list(body.xpath(".//script | .//style | .//noscript | .//nav | .//footer | .//form | .//svg")):
            if node.getparent() is not None:
                node.drop_tree()
        text = "\n".join(piece.strip() for piece in body.itertext() if piece.strip())
        if len(text) < 160:
            result["extraction_status"] = "needs-browser-rendering-or-content-review"
        else:
            result["units"] = [{"locator": "html:main" if mains else "html:document", "text": text}]
            result["extraction_status"] = "extracted-unreviewed"
        result["content_class"] = "html-page-completeness-unreviewed"
        return result
    mime = content_type.split(";", 1)[0].strip().lower()
    if mime.startswith("text/") or mime in {"application/json", "application/xml"} or mime.endswith(("+json", "+xml")):
        text = data.decode("utf-8", errors="replace").strip()
        result["units"] = [{"locator": "document", "text": text}] if text else []
        result["extraction_status"] = "extracted-unreviewed" if text else "empty-document"
        result["content_class"] = "text-document"
    return result


def attach_extraction(record: dict, data: bytes, output_root: Path) -> dict:
    for key in ("text_path", "text_sha256", "extraction_error"):
        record.pop(key, None)
    try:
        extracted = extract(data, record["content_type"], record["response_url"])
    except Exception as exc:
        extracted = {"title": "", "units": [], "links": [], "extraction_status": "extraction-error",
                     "extraction_error": f"{type(exc).__name__}: {exc}", "content_class": "unclassified"}
    units = extracted.pop("units")
    record.update(extracted)
    record["text_character_count"] = sum(len(unit["text"]) for unit in units)
    record["extraction_revision"] = 4
    if units:
        rendered = (json.dumps(units, ensure_ascii=True, indent=2) + "\n").encode("utf-8")
        text_path, text_digest = save_object(output_root, rendered, ".json")
        record.update({"text_path": text_path, "text_sha256": text_digest})
    return record


class SafeRedirect(HTTPRedirectHandler):
    def __init__(self, hosts: set[str]):
        self.hosts = hosts

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if not allowed_url(newurl, self.hosts):
            raise ValueError("Redirect requires separate authority/HTTPS review")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


class Fetcher:
    def __init__(self, hosts: set[str], timeout: int, max_bytes: int, delay: float):
        self.hosts, self.timeout, self.max_bytes, self.delay = hosts, timeout, max_bytes, delay
        self.locks = defaultdict(threading.Lock)
        self.last = defaultdict(float)
        self.robots: dict[str, tuple[RobotFileParser | None, str | None]] = {}
        self.paused: dict[str, str] = {}
        self.denials = Counter()

    def request(self, url: str, max_bytes: int) -> tuple[bytes, dict, str, int]:
        if not allowed_url(url, self.hosts):
            raise ValueError("Source requires authority/HTTPS review")
        opener = build_opener(SafeRedirect(self.hosts))
        with opener.open(Request(transport_url(url), headers={"User-Agent": USER_AGENT}), timeout=self.timeout) as response:
            chunks, size, started = [], 0, time.monotonic()
            while True:
                if time.monotonic() - started > max(60, self.timeout * 4):
                    raise TimeoutError("Total response-read time exceeded; preserve as an acquisition gap")
                chunk = response.read1(min(65536, max_bytes + 1 - size))
                if not chunk:
                    break
                chunks.append(chunk)
                size += len(chunk)
                if size > max_bytes:
                    raise ValueError("Source exceeds configured byte limit; use a reviewed large-file download")
            data = b"".join(chunks)
            headers = {key.lower(): value for key, value in response.headers.items()
                       if key.lower() in {"content-type", "content-length", "etag", "last-modified"}}
            return data, headers, response.geturl(), response.status

    def permission(self, url: str) -> tuple[bool, str]:
        host = urlparse(url).netloc
        if host not in self.robots:
            robots_url = f"https://{host}/robots.txt"
            try:
                data, _, _, _ = self.request(robots_url, 1024 * 1024)
                parser = RobotFileParser()
                parser.parse(data.decode("utf-8", errors="replace").splitlines())
                self.robots[host] = parser, None
            except HTTPError as exc:
                self.robots[host] = (None, None if exc.code in (404, 410) else f"robots-http-{exc.code}")
            except Exception as exc:
                self.robots[host] = None, f"robots-check-failed:{type(exc).__name__}"
        parser, error = self.robots[host]
        if error:
            return False, error
        if parser and not parser.can_fetch(USER_AGENT, url):
            return False, "robots-disallowed"
        return True, "robots-allowed" if parser else "robots-not-published"

    def fetch(self, item: dict, output_root: Path) -> dict:
        url = item["canonical_url"]
        record = {**item, "attempted_at": now(), "acquisition_method": "https-original-bytes",
                  "legal_review_status": "not-reviewed", "current_law_release": False,
                  "redistribution_status": "not-cleared", "source_time_version": "not-established"}
        host = urlparse(url).netloc
        try:
            with self.locks[host]:
                if host in self.paused:
                    return {**record, "capture_status": "deferred", "reason": self.paused[host]}
                allowed, reason = self.permission(url)
                if not allowed:
                    return {**record, "capture_status": "deferred", "reason": reason}
                robot = self.robots[host][0]
                delay = max(self.delay, (robot.crawl_delay(USER_AGENT) or robot.crawl_delay("*") or 0) if robot else 0)
                time.sleep(max(0, delay - (time.monotonic() - self.last[host])))
                self.last[host] = time.monotonic()
                data, headers, response_url, status = self.request(url, self.max_bytes)
                self.denials[host] = 0
            content_type = headers.get("content-type", "application/octet-stream")
            suffix = ".pdf" if data.startswith(b"%PDF-") else ".html" if "html" in content_type else ".bin"
            snapshot_path, digest = save_object(output_root, data, suffix)
            record.update({"capture_status": "bytes-preserved", "retrieved_at": now(),
                           "response_url": response_url, "status_code": status, "response_headers": headers,
                           "content_type": content_type, "size_bytes": len(data), "snapshot_path": snapshot_path,
                           "sha256": digest, "robots_status": reason})
            return attach_extraction(record, data, output_root)
        except HTTPError as exc:
            reason = f"http-{exc.code}"
            if exc.code in (401, 403, 429):
                with self.locks[host]:
                    self.denials[host] += 1
                    if exc.code == 429 or self.denials[host] >= 3:
                        self.paused[host] = f"host-paused-after-{reason}; review access without bypassing restrictions"
            return {**record, "capture_status": "failed", "reason": reason}
        except Exception as exc:
            return {**record, "capture_status": "failed", "reason": f"{type(exc).__name__}: {exc}"}


def discover(row: dict, hosts: set[str], max_depth: int) -> list[dict]:
    if row.get("extraction_status") not in {"extracted-unreviewed", "needs-browser-rendering-or-content-review"} or row["discovery_depth"] >= max_depth:
        return []
    results = []
    for link in row.get("links", []):
        url = link["url"]
        if not ATTACHMENT.search(url) or not allowed_url(url, hosts):
            continue
        label = link.get("label") or link.get("text", "")
        if NON_DOCUMENT.search(label) and not re.search(r"annual.report|compliance|electric|energy", label, re.I):
            continue
        results.append({"canonical_url": url, "references": [f"attachment:{row['canonical_url']}"],
                        "source_family_ids": row["source_family_ids"],
                        "discovery_depth": row["discovery_depth"] + 1,
                        "discovered_from": [row["canonical_url"]], "link_label": label,
                        "relevance_status": "linked-attachment-needs-review"})
    return results


def balanced_queue(items: list[dict]) -> deque:
    domains: dict[str, deque] = defaultdict(deque)
    for item in items:
        domains[host_key(item["canonical_url"])].append(item)
    queue: deque = deque()
    while domains:
        for host in list(domains):
            queue.append(domains[host].popleft())
            if not domains[host]:
                del domains[host]
    return queue


def take_available_host(queue: deque, active_items: list[dict]) -> dict | None:
    busy = {host_key(item["canonical_url"]) for item in active_items}
    for _ in range(len(queue)):
        item = queue.popleft()
        if host_key(item["canonical_url"]) not in busy:
            return item
        queue.append(item)
    return None


def generic_reextract_allowed(row: dict) -> bool:
    return (row.get("capture_status") == "bytes-preserved"
            and not row.get("ocr_pages_path") and not row.get("embedded_record_path")
            and not row.get("browser_capture_provenance"))


def rebuild_outputs(root: Path, entries: dict[str, dict], seed_count: int, run_id: str, legacy_root: Path | None = None) -> dict:
    rows = read_jsonl(root / "source-originals" / "manifest.jsonl")
    latest, preserved = select_manifest_rows(rows)
    inventory = [{**item, "latest_capture_status": latest.get(url, {}).get("capture_status", "not-attempted"),
                  "latest_reason": latest.get(url, {}).get("reason"),
                  "latest_extraction_status": latest.get(url, {}).get("extraction_status"),
                  "has_research_text": url in preserved}
                 for url, item in sorted(entries.items())]
    data_root = root / "data"
    data_root.mkdir(parents=True, exist_ok=True)
    (data_root / "source-original-inventory.jsonl").write_text(
        "".join(json.dumps(row, ensure_ascii=True) + "\n" for row in inventory), encoding="utf-8")
    gaps = []
    for item in inventory:
        source = latest.get(item["canonical_url"], {})
        tasks = []
        if item.get("selection_status") == "authority-or-https-review-required":
            tasks.append("authority-or-https-review")
        elif item["latest_capture_status"] == "not-attempted":
            tasks.append("pending-acquisition")
        elif source.get("capture_status") in {"failed", "deferred"}:
            tasks.append("access-or-source-location-review")
        if source.get("extraction_status") in {"needs-browser-rendering-or-content-review", "blocked-or-error-page"}:
            tasks.append("permitted-browser-or-source-review")
        if source.get("extraction_status") in {"needs-ocr-or-blank-page-review", "ocr-extracted-unreviewed"}:
            tasks.append("ocr-or-blank-page-review")
        if source.get("extraction_status") in {"unsupported-format", "extraction-error"}:
            tasks.append("format-specific-extraction-review")
        if item["has_research_text"]:
            tasks.append("relevance-completeness-temporal-and-rights-review")
        gaps.append({"canonical_url": item["canonical_url"], "tasks": tasks,
                     "reason": item["latest_reason"], "source_family_ids": item["source_family_ids"],
                     "pages_without_text": source.get("pages_without_text", []),
                     "current_law_release": False})
    (data_root / "source-original-gap-queue.jsonl").write_text(
        "".join(json.dumps(row, ensure_ascii=True) + "\n" for row in gaps), encoding="utf-8")
    database = root / "source-originals" / "search.sqlite3"
    with closing(sqlite3.connect(database)) as connection, connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("DROP TABLE IF EXISTS originals")
        connection.execute("CREATE VIRTUAL TABLE originals USING fts5(url UNINDEXED, title, locator UNINDEXED, text, sha256 UNINDEXED, snapshot_path UNINDEXED, review_status UNINDEXED, retrieved_at UNINDEXED, acquisition_method UNINDEXED, extraction_status UNINDEXED, page_count UNINDEXED, pages_without_text UNINDEXED, text_path UNINDEXED, text_sha256 UNINDEXED)")
        for url, row in preserved.items():
            text_path = root / row["text_path"]
            raw = text_path.read_bytes()
            if hashlib.sha256(raw).hexdigest() != row["text_sha256"]:
                raise ValueError(f"Text object failed integrity validation: {text_path}")
            for unit in json.loads(raw):
                connection.execute("INSERT INTO originals VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (
                    url, row.get("title", ""), unit["locator"], unit["text"], row["sha256"],
                    row["snapshot_path"], "research-only; not legal clearance", row.get("retrieved_at"),
                    row.get("acquisition_method", "https-original-bytes"), row.get("extraction_status"),
                    row.get("page_count"), json.dumps(row.get("pages_without_text", [])),
                    row["text_path"], row["text_sha256"]))
    summary = {"schema_version": "1.0", "generated_at": now(), "run_id": run_id,
               "initial_seed_url_count": seed_count, "discovered_attachment_url_count": len(entries) - seed_count,
               "inventory_url_count": len(entries), "attempted_distinct_url_count": len(latest),
               "response_bytes_preserved_url_count": sum(row["capture_status"] == "bytes-preserved" for row in latest.values()),
               "browser_text_preserved_url_count": sum(row["capture_status"] == "browser-text-preserved" for row in latest.values()),
               "research_text_url_count": len(preserved), "legally_verified_full_text_url_count": 0,
               "not_attempted_url_count": sum(row["latest_capture_status"] == "not-attempted" for row in inventory),
               "latest_capture_status_counts": dict(Counter(row["capture_status"] for row in latest.values())),
               "latest_extraction_status_counts": dict(Counter(row.get("extraction_status", "not-extracted") for row in latest.values())),
               "failure_reason_counts": dict(Counter(row["reason"] for row in latest.values() if row.get("reason"))),
               "public_exhaustiveness_status": "not-established; source and period reconciliation still required",
               "legacy_answer_index_modified": False,
               "limitation": "Preserved response bytes can include an error or application shell. Extracted research text is not proof of legal currency, complete page/attachment coverage, source rights or semantic correctness."}
    if legacy_root:
        remote_urls = {row["canonical_url"] for row in read_jsonl(legacy_root / "data" / "source-artifact-ledger.jsonl")
                       if row.get("snapshot_status") == "remote-only-no-snapshot"}
        summary["legacy_remote_url_count"] = len(remote_urls)
        summary["legacy_remote_urls_with_research_text"] = sum(normalise_url(url) in preserved for url in remote_urls)
        summary["legacy_remote_urls_still_without_research_text"] = sum(normalise_url(url) not in preserved for url in remote_urls)
        summary["legacy_fragment_aliases_with_research_text"] = sum(url not in preserved and normalise_url(url) in preserved for url in remote_urls)
    (data_root / "source-original-summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    report_root = root / "review" / "results"
    report_root.mkdir(parents=True, exist_ok=True)
    report = ["# Original Source Collection Report", "", f"Generated: {summary['generated_at']}", "",
              "This is acquisition evidence, not a claim of complete Australian coverage or verified legal answers.", "",
              "| Measure | Count |", "|---|---:|",
              f"| Initial cited seed URLs | {seed_count} |",
              f"| Discovered official attachment candidates | {len(entries) - seed_count} |",
              f"| Total known URL inventory | {len(entries)} |",
              f"| URLs with recorded acquisition outcomes | {len(latest)} |",
              f"| Latest HTTP responses preserved, including shells/errors | {summary['response_bytes_preserved_url_count']} |",
              f"| Browser-rendered text copies preserved | {summary['browser_text_preserved_url_count']} |",
              f"| URLs with searchable, unreviewed research text | {len(preserved)} |",
              f"| URLs not yet attempted, including authority-review candidates | {summary['not_attempted_url_count']} |",
              "| Full-text sources certified for operative legal use in this layer | 0 |", "",
              "## Follow-up queues", ""]
    for task, count in sorted(Counter(task for row in gaps for task in row["tasks"]).items()):
        report.append(f"- {task}: {count}")
    if legacy_root:
        report.extend(["", "## Legacy snapshot-gap reconciliation", "",
                       f"Of {summary['legacy_remote_url_count']} legacy remote-only URLs, {summary['legacy_remote_urls_with_research_text']} now have research text in the new archive; {summary['legacy_remote_urls_still_without_research_text']} still do not.",
                       "The legacy approved-answer index and source ledger were not relabelled as complete or current."])
    report.extend(["", "## Publisher coverage of known URLs only", "",
                   "This table reconciles the discovered inventory, not each publisher's full historical holdings.", "",
                   "| Publisher host | Known URLs | With research text | Without research text |",
                   "|---|---:|---:|---:|"])
    host_totals = Counter(host_key(item["canonical_url"]) for item in inventory)
    host_text = Counter(host_key(item["canonical_url"]) for item in inventory if item["has_research_text"])
    for host, count in sorted(host_totals.items()):
        report.append(f"| {host} | {count} | {host_text[host]} | {count - host_text[host]} |")
    report.extend(["", "The URL inventory is not an established universe of all Australian electricity material. Site-specific pagination, historical versions, missing source families, inaccessible material and relevance review remain open.",
                   "Source responses, extracted text and browser renderings are different representations. Preserve their individual provenance and do not mistake a current capture for an event-time copy.",
                   "Raw files and derived full text are excluded from Git until redistribution rights are reviewed.", ""])
    (report_root / "original-source-collection-report.md").write_text("\n".join(report), encoding="utf-8")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--max-requests", type=int, default=1000)
    parser.add_argument("--max-runtime-seconds", type=int, default=900, help="Stop scheduling after this budget, then finish in-flight requests.")
    parser.add_argument("--workers", type=int, choices=range(1, 5), default=4)
    parser.add_argument("--timeout", type=int, default=15)
    parser.add_argument("--max-mb", type=int, default=40)
    parser.add_argument("--delay", type=float, default=1.0)
    parser.add_argument("--attachment-depth", type=int, choices=range(0, 3), default=1)
    parser.add_argument("--inventory-only", action="store_true")
    parser.add_argument("--retry-failed", action="store_true")
    parser.add_argument("--resume-host-deferred", action="store_true", help="Process never-requested URLs deferred after another URL's 401/403; do not retry the denied URL.")
    parser.add_argument("--reextract", action="store_true", help="Re-extract preserved bytes offline; requires --inventory-only.")
    parser.add_argument("--host", action="append", default=[])
    args = parser.parse_args()
    if args.max_requests < 0 or args.timeout < 1 or args.max_mb < 1 or args.delay < 0 or args.max_runtime_seconds < 1:
        parser.error("Invalid resource limit")
    if args.reextract and not args.inventory_only:
        parser.error("--reextract requires --inventory-only to keep offline processing separate from fetching")
    source_root = args.root.resolve()
    output_root = (args.output_root or source_root).resolve()
    entries, hosts = seeds(source_root)
    seed_count = len(entries)
    for item in read_jsonl(output_root / "data" / "source-original-inventory.jsonl"):
        entries.setdefault(item["canonical_url"], {key: value for key, value in item.items()
                                                  if not key.startswith("latest_") and key != "has_research_text"})
    run_id = now()
    originals = output_root / "source-originals"
    originals.mkdir(parents=True, exist_ok=True)
    manifest = originals / "manifest.jsonl"
    previous = {row["canonical_url"]: row for row in read_jsonl(manifest)}
    if args.reextract:
        with manifest.open("a", encoding="utf-8", newline="\n") as handle:
            for row in previous.values():
                if not generic_reextract_allowed(row) or row.get("extraction_revision") == 4:
                    continue
                data = (output_root / row["snapshot_path"]).read_bytes()
                if hashlib.sha256(data).hexdigest() != row["sha256"]:
                    raise ValueError("Preserved source hash mismatch before re-extraction")
                attach_extraction(row, data, output_root)
                row["invalidates_prior_text"] = True
                row["extraction_refreshed_at"] = now()
                handle.write(json.dumps(row, ensure_ascii=True) + "\n")
                handle.flush()
    # Recover attachment discoveries even when an interrupted run has no rebuilt inventory.
    for row in previous.values():
        entries.setdefault(row["canonical_url"], {key: row[key] for key in
                           ("canonical_url", "references", "source_family_ids", "discovery_depth", "discovered_from")})
        for child in discover(row, hosts, args.attachment_depth):
            entries.setdefault(child["canonical_url"], child)
    candidates = []
    for url, item in entries.items():
        item["selection_status"] = "public-authority-candidate" if allowed_url(url, hosts) else "authority-or-https-review-required"
        if not allowed_url(url, hosts) or (args.host and host_key(url) not in args.host):
            continue
        if url in previous:
            prior = previous[url]
            host_deferred = (args.resume_host_deferred and prior["capture_status"] == "deferred"
                             and re.match(r"host-paused-after-http-(401|403);", prior.get("reason", "")))
            if not host_deferred and (not args.retry_failed or prior["capture_status"] in {"bytes-preserved", "browser-text-preserved"}):
                continue
        candidates.append(item)
    print(json.dumps({"seed_urls": seed_count, "queued_urls": len(candidates), "max_requests": args.max_requests}), flush=True)
    queue = balanced_queue(candidates)
    fetcher = Fetcher(hosts, args.timeout, args.max_mb * 1024 * 1024, args.delay)
    attempted = 0
    deadline = time.monotonic() + args.max_runtime_seconds
    if not args.inventory_only:
        with manifest.open("a", encoding="utf-8", newline="\n") as handle, ThreadPoolExecutor(max_workers=args.workers) as pool:
            active = {}
            while active or (queue and attempted < args.max_requests and time.monotonic() < deadline):
                while queue and len(active) < args.workers and attempted < args.max_requests and time.monotonic() < deadline:
                    item = take_available_host(queue, list(active.values()))
                    if item is None:
                        break
                    active[pool.submit(fetcher.fetch, item, output_root)] = item
                    attempted += 1
                if not active:
                    break
                done, _ = wait(active, return_when=FIRST_COMPLETED)
                for future in done:
                    active.pop(future)
                    row = future.result()
                    row["run_id"] = run_id
                    handle.write(json.dumps(row, ensure_ascii=True) + "\n")
                    handle.flush()
                    for child in discover(row, hosts, args.attachment_depth):
                        url = child["canonical_url"]
                        if url not in entries:
                            entries[url] = child
                            if not args.host or host_key(url) in args.host:
                                queue.append(child)
                    completed = attempted - len(active)
                    if completed % 25 == 0:
                        print(f"Checkpoint: {completed} URL outcomes recorded; {len(queue)} queued", flush=True)
    summary = rebuild_outputs(output_root, entries, seed_count, run_id, source_root)
    print(json.dumps(summary, indent=2, ensure_ascii=True), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
